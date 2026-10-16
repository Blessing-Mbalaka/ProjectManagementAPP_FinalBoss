from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from adminpanel.models import ResearchCentre
from projects.models import Project
from projects.workbook_import.schema import PROJECT_COLUMNS, TASK_COLUMNS, WorkbookReadResult
from projects.workbook_import.tokens import PreviewTokenError, sign_preview, unsign_preview
from projects.workbook_import.validation import validate_workbook
from projects.workbook_import.workbook import build_template_workbook


User = get_user_model()
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ProjectWorkbookUITests(TestCase):
    def setUp(self):
        self.centre = ResearchCentre.objects.create(name="UI Centre")
        self.manager = User.objects.create_user(username="manager1", password="pass", role="manager", research_centre=self.centre)
        self.alice = User.objects.create_user(username="alice", role="staff", research_centre=self.centre)
        self.client = Client()

    def workbook_upload(self, *, import_key="ui-project-1", project_name="UI Project", parent="Parent"):
        workbook = load_workbook(BytesIO(build_template_workbook()))
        projects = workbook["Projects"]
        projects.delete_rows(2, projects.max_row)
        row = {column: "" for column in PROJECT_COLUMNS}
        row.update({
            "import_key": import_key,
            "project_name": project_name,
            "portfolio": "UI Portfolio",
            "description": "Uploaded from the UI",
            "project_type": "software",
            "project_status": "planning",
            "due_date": "2027-03-31",
            "project_manager_username": "manager1",
            "research_centre": self.centre.name,
            "team_usernames": "alice",
        })
        projects.append([row[column] for column in PROJECT_COLUMNS])
        tasks = workbook["Tasks"]
        tasks.delete_rows(2, tasks.max_row)
        parent_row = {column: "" for column in TASK_COLUMNS}
        parent_row.update({"project_name": project_name, "task_name": "Parent", "status": "todo", "priority": "High", "task_type": "Other"})
        child_row = {column: "" for column in TASK_COLUMNS}
        child_row.update({"project_name": project_name, "task_name": "Child", "parent_task": parent, "assigned_to": "alice", "status": "todo", "priority": "Medium", "task_type": "Other"})
        tasks.append([parent_row[column] for column in TASK_COLUMNS])
        tasks.append([child_row[column] for column in TASK_COLUMNS])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile("projects.xlsx", output.getvalue(), content_type=XLSX_MIME)

    def test_download_requires_an_authorized_project_management_role(self):
        response = self.client.get(reverse("download_project_workbook"))
        self.assertEqual(response.status_code, 302)

        for role in ("staff", "student"):
            user = User.objects.create_user(username=role, role=role, research_centre=self.centre)
            self.client.force_login(user)
            self.assertEqual(self.client.get(reverse("download_project_workbook")).status_code, 403)

        for role in ("manager", "centrehead", "dean", "admin"):
            user = self.manager if role == "manager" else User.objects.create_user(username=f"allowed-{role}", role=role, research_centre=self.centre)
            self.client.force_login(user)
            response = self.client.get(reverse("download_project_workbook"))
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["Content-Type"], XLSX_MIME)
            self.assertIn("project-import-template.xlsx", response["Content-Disposition"])

    def test_valid_preview_shows_changes_and_confirmation(self):
        self.client.force_login(self.manager)

        response = self.client.post(reverse("preview_project_workbook"), {"workbook": self.workbook_upload()})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "UI Project")
        self.assertContains(response, "UI Portfolio")
        self.assertContains(response, "Parent")
        self.assertContains(response, "Child")
        self.assertContains(response, 'name="preview_token"')
        self.assertTrue(response.context["preview"].is_valid)

    def test_invalid_parent_shows_row_error_without_confirmation(self):
        self.client.force_login(self.manager)

        response = self.client.post(reverse("preview_project_workbook"), {"workbook": self.workbook_upload(parent="Missing parent")})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Tasks")
        self.assertContains(response, "Missing parent")
        self.assertNotContains(response, 'name="preview_token"')
        self.assertFalse(Project.objects.exists())

    def test_preview_tokens_reject_tampering_expiry_and_another_user(self):
        read_result = WorkbookReadResult(project_rows=(), task_rows=())
        preview = validate_workbook(read_result, actor=self.manager)
        with patch("django.core.signing.time.time", return_value=100):
            token = sign_preview(preview, actor_id=self.manager.id)

        with self.assertRaises(PreviewTokenError):
            unsign_preview(token + "tampered", actor_id=self.manager.id)
        with patch("django.core.signing.time.time", return_value=2000):
            with self.assertRaises(PreviewTokenError):
                unsign_preview(token, actor_id=self.manager.id, max_age=1800)
        with self.assertRaises(PermissionError):
            unsign_preview(token, actor_id=self.alice.id, max_age=None)

    def test_confirmation_applies_idempotently_and_rechecks_scope(self):
        self.client.force_login(self.manager)
        preview_response = self.client.post(reverse("preview_project_workbook"), {"workbook": self.workbook_upload()})
        token = preview_response.context["preview_token"]

        response = self.client.post(reverse("apply_project_workbook"), {"preview_token": token})
        replay = self.client.post(reverse("apply_project_workbook"), {"preview_token": token})

        self.assertRedirects(response, reverse("assign_projects"))
        self.assertRedirects(replay, reverse("assign_projects"))
        self.assertEqual(Project.objects.count(), 1)
        self.assertEqual(Project.objects.get().tasks.count(), 2)

        owned = Project.objects.get()
        other_manager = User.objects.create_user(username="manager2", role="manager", research_centre=self.centre)
        owned.created_by = other_manager
        owned.assigned_user = other_manager
        owned.save(update_fields=["created_by", "assigned_user"])
        preview_response = self.client.post(reverse("preview_project_workbook"), {
            "workbook": self.workbook_upload(import_key=owned.import_key, project_name=owned.name),
        })
        self.assertEqual(preview_response.status_code, 200)
        self.assertFalse(preview_response.context["preview"].is_valid)

    def test_apply_rejects_token_owned_by_another_user(self):
        self.client.force_login(self.manager)
        preview_response = self.client.post(reverse("preview_project_workbook"), {"workbook": self.workbook_upload()})
        token = preview_response.context["preview_token"]
        other_manager = User.objects.create_user(username="manager2", password="pass", role="manager", research_centre=self.centre)
        self.client.force_login(other_manager)

        response = self.client.post(reverse("apply_project_workbook"), {"preview_token": token})

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Project.objects.exists())
