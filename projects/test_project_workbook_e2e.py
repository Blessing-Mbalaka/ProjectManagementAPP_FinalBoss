from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from adminpanel.models import ResearchCentre
from manager.models import Book
from projects.models import Assignment, Project, Task
from projects.workbook_import.application import apply_workbook_preview
from projects.workbook_import.schema import PROJECT_COLUMNS, TASK_COLUMNS
from projects.workbook_import.validation import validate_workbook
from projects.workbook_import.workbook import XLSX_MIME, build_template_workbook, read_workbook


User = get_user_model()


class ProjectWorkbookEndToEndTests(TestCase):
    def setUp(self):
        self.centre = ResearchCentre.objects.create(name="End-to-End Centre")
        self.manager = User.objects.create_user(
            username="manager1",
            password="pass",
            role="manager",
            research_centre=self.centre,
        )
        self.lead = User.objects.create_user(username="lead", role="staff", research_centre=self.centre)
        self.coauthor = User.objects.create_user(username="coauthor", role="staff", research_centre=self.centre)

    def workbook(self):
        workbook = load_workbook(BytesIO(build_template_workbook()))
        projects = workbook["Projects"]
        projects.delete_rows(2, projects.max_row)
        rows = [
            {
                "import_key": "e2e-software",
                "project_name": "E2E Software",
                "portfolio": "Digital",
                "description": "Software project",
                "project_type": "software",
                "project_status": "planning",
                "project_manager_username": "manager1",
                "research_centre": self.centre.name,
                "team_usernames": "lead;coauthor",
            },
            {
                "import_key": "e2e-book",
                "project_name": "E2E Book",
                "portfolio": "Publications",
                "description": "Book project",
                "project_type": "book",
                "project_status": "in-progress",
                "project_manager_username": "manager1",
                "research_centre": self.centre.name,
                "publication_status": "writing",
                "lead_author_username": "lead",
                "co_author_usernames": "coauthor",
                "external_authors": "External Scholar",
                "publisher": "Research Press",
                "total_chapters": "3",
                "completed_chapters": "1",
            },
        ]
        for values in rows:
            row = {column: "" for column in PROJECT_COLUMNS}
            row.update(values)
            projects.append([row[column] for column in PROJECT_COLUMNS])

        tasks = workbook["Tasks"]
        tasks.delete_rows(2, tasks.max_row)
        task_rows = [
            {"project_name": "E2E Software", "task_name": "Build", "assigned_to": "lead", "status": "in_progress", "priority": "High", "task_type": "Backend"},
            {"project_name": "E2E Software", "task_name": "Test", "parent_task": "Build", "assigned_to": "coauthor", "status": "todo", "priority": "Medium", "task_type": "Testing"},
        ]
        for values in task_rows:
            row = {column: "" for column in TASK_COLUMNS}
            row.update(values)
            tasks.append([row[column] for column in TASK_COLUMNS])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_workbook_flows_through_services_and_existing_manager_pages(self):
        uploaded = SimpleUploadedFile("projects.xlsx", self.workbook(), content_type=XLSX_MIME)
        preview = validate_workbook(read_workbook(uploaded), actor=self.manager)

        result = apply_workbook_preview(preview, actor=self.manager)

        self.assertTrue(preview.is_valid, preview.errors)
        self.assertEqual(result.created_projects, 2)
        book = Book.objects.select_related("project", "lead_author_user").get()
        self.assertEqual(book.lead_author_user, self.lead)
        self.assertCountEqual(book.co_authors_users.all(), [self.coauthor])
        self.assertEqual(book.external_authors, ["External Scholar"])
        software = Project.objects.get(import_key="e2e-software")
        self.assertEqual(software.portfolio, "Digital")
        self.assertEqual(software.tasks.get(title="Test").parent_task.title, "Build")
        self.assertTrue(Assignment.objects.filter(project=software, team_member__user=self.coauthor).exists())

        client = Client()
        client.force_login(self.manager)
        assignment_page = client.get(reverse("assign_projects"), {"project_id": software.id})
        kanban_page = client.get(reverse("manager_kanban"))
        book_page = client.get(reverse("manager_book"))

        self.assertContains(assignment_page, "E2E Software")
        self.assertContains(assignment_page, "Build")
        self.assertContains(assignment_page, "Test")
        self.assertContains(kanban_page, "E2E Software")
        self.assertContains(book_page, "E2E Book")
        self.assertEqual(Task.objects.count(), 2)
