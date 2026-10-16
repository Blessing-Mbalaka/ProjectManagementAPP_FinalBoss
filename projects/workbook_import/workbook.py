from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from .schema import (
    MAX_PROJECT_ROWS,
    MAX_TASK_ROWS,
    MAX_UPLOAD_BYTES,
    PROJECT_COLUMNS,
    SHEET_NAMES,
    TASK_COLUMNS,
    WorkbookCellError,
    WorkbookReadResult,
)


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _style_header(sheet):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="580658")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def build_template_workbook() -> bytes:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instruction_rows = [
        ("Project workbook import", "Complete Projects before Tasks."),
        ("Dates", "Use YYYY-MM-DD."),
        ("Multiple values", "Separate usernames or external authors with a semicolon (;)."),
        ("Project types", "software; book; paper; conference"),
        ("Project statuses", "planning; in-progress; on-hold; completed"),
        ("Task statuses", "todo; in_progress; review; done"),
        ("Priorities", "Low; Medium; High"),
        ("Hierarchy", "Leave parent_task blank for a main task; otherwise use the exact task_name in the same project."),
        ("Authors", "Use usernames for system authors and external_authors only for people without accounts."),
        ("Re-imports", "import_key identifies projects; normalized task names identify tasks within a project."),
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    instructions["A1"].font = Font(bold=True, color="FFFFFF")
    instructions["B1"].font = Font(bold=True, color="FFFFFF")
    instructions["A1"].fill = PatternFill("solid", fgColor="580658")
    instructions["B1"].fill = PatternFill("solid", fgColor="580658")

    projects = workbook.create_sheet("Projects")
    projects.append(PROJECT_COLUMNS)
    project_example = {column: "" for column in PROJECT_COLUMNS}
    project_example.update({
        "import_key": "website-redesign-001",
        "project_name": "Website Redesign",
        "portfolio": "Digital Projects",
        "description": "Refresh the public website",
        "project_type": "software",
        "project_status": "planning",
        "due_date": "2027-03-31",
        "project_manager_username": "manager_username",
        "research_centre": "Research Centre Name",
        "team_usernames": "alice;bob",
    })
    projects.append([project_example[column] for column in PROJECT_COLUMNS])
    _style_header(projects)

    tasks = workbook.create_sheet("Tasks")
    tasks.append(TASK_COLUMNS)
    task_examples = [
        ("Website Redesign", "Build frontend", "", "alice", "Frontend lead", "todo", "High", "Frontend", "2027-02-28", ""),
        ("Website Redesign", "Create homepage", "Build frontend", "bob", "Homepage", "todo", "Medium", "Frontend", "2027-02-15", ""),
        ("Website Redesign", "Create contact page", "Build frontend", "alice", "Contact page", "todo", "Medium", "Frontend", "2027-02-20", ""),
    ]
    for row in task_examples:
        tasks.append(row)
    _style_header(tasks)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _error(message, *, sheet="Workbook", row=None, field="file"):
    return WorkbookCellError(sheet=sheet, row=row, field=field, message=message)


def _read_rows(sheet, columns, limit):
    errors = []
    header = tuple((cell.value or "") for cell in sheet[1])
    if header != columns:
        errors.append(_error(
            f"Expected headers: {', '.join(columns)}.",
            sheet=sheet.title,
            row=1,
            field="headers",
        ))
        return (), errors

    rows = []
    for excel_row, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        if any(cell.data_type == "f" for cell in cells):
            errors.append(_error(
                "Formula cells are not allowed.",
                sheet=sheet.title,
                row=excel_row,
                field="formula",
            ))
            continue
        values = [cell.value for cell in cells[:len(columns)]]
        if not any(value not in (None, "") for value in values):
            continue
        row = dict(zip(columns, values))
        row["_row"] = excel_row
        rows.append(row)
    if len(rows) > limit:
        label = f"{limit:,}"
        errors.append(_error(
            f"{sheet.title} may contain at most {label} data rows.",
            sheet=sheet.title,
            field="rows",
        ))
    return tuple(rows[:limit]), errors


def read_workbook(upload, *, max_bytes=MAX_UPLOAD_BYTES) -> WorkbookReadResult:
    name = str(getattr(upload, "name", ""))
    if not name.lower().endswith(".xlsx"):
        return WorkbookReadResult(errors=(_error("Upload an .xlsx workbook."),))
    size = getattr(upload, "size", None)
    if size is not None and size > max_bytes:
        return WorkbookReadResult(errors=(_error("Workbook must not exceed 2 MiB."),))
    data = upload.read()
    if len(data) > max_bytes:
        return WorkbookReadResult(errors=(_error("Workbook must not exceed 2 MiB."),))

    try:
        with ZipFile(BytesIO(data), "r") as archive:
            if any(path.lower().endswith("vbaproject.bin") for path in archive.namelist()):
                return WorkbookReadResult(errors=(_error("Macro-enabled workbooks are not allowed."),))
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_vba=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        return WorkbookReadResult(errors=(_error("The uploaded file is not a valid .xlsx workbook."),))

    errors = []
    if tuple(workbook.sheetnames) != SHEET_NAMES:
        missing = [name for name in SHEET_NAMES if name not in workbook.sheetnames]
        extra = [name for name in workbook.sheetnames if name not in SHEET_NAMES]
        if missing:
            errors.append(_error(f"Missing required sheet(s): {', '.join(missing)}.", field="sheets"))
        if extra:
            errors.append(_error(f"Unexpected sheet(s): {', '.join(extra)}.", field="sheets"))
    if errors:
        return WorkbookReadResult(errors=tuple(errors))

    project_rows, project_errors = _read_rows(workbook["Projects"], PROJECT_COLUMNS, MAX_PROJECT_ROWS)
    task_rows, task_errors = _read_rows(workbook["Tasks"], TASK_COLUMNS, MAX_TASK_ROWS)
    return WorkbookReadResult(
        project_rows=project_rows,
        task_rows=task_rows,
        errors=tuple(project_errors + task_errors),
    )
