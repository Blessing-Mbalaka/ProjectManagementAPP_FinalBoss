from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from projects.workbook_import.schema import PROJECT_COLUMNS, TASK_COLUMNS
from projects.workbook_import.workbook import build_template_workbook, read_workbook


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(project_rows=(), task_rows=(), project_headers=None, task_headers=None):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    projects = workbook.create_sheet("Projects")
    tasks = workbook.create_sheet("Tasks")
    projects.append(project_headers or PROJECT_COLUMNS)
    tasks.append(task_headers or TASK_COLUMNS)
    for row in project_rows:
        projects.append(row)
    for row in task_rows:
        tasks.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload(data, name="projects.xlsx"):
    return SimpleUploadedFile(name, data, content_type=XLSX_MIME)


class WorkbookTemplateTests(SimpleTestCase):
    def test_template_has_exact_contract_and_hierarchy_example(self):
        workbook = load_workbook(BytesIO(build_template_workbook()), read_only=True, data_only=False)

        self.assertEqual(workbook.sheetnames, ["Instructions", "Projects", "Tasks"])
        self.assertEqual(tuple(cell.value for cell in workbook["Projects"][1]), PROJECT_COLUMNS)
        self.assertEqual(tuple(cell.value for cell in workbook["Tasks"][1]), TASK_COLUMNS)
        task_rows = list(workbook["Tasks"].iter_rows(min_row=2, values_only=True))
        self.assertIn("Website Redesign", [row[0] for row in task_rows])
        self.assertIn("Build frontend", [row[2] for row in task_rows])
        instructions = "\n".join(str(cell.value or "") for row in workbook["Instructions"] for cell in row)
        self.assertIn("YYYY-MM-DD", instructions)
        self.assertIn("semicolon", instructions.lower())
        self.assertIn("software", instructions)
        self.assertIn("conference", instructions)


class WorkbookReaderSafetyTests(SimpleTestCase):
    def test_missing_sheet_and_wrong_headers_are_reported(self):
        missing = Workbook()
        missing.active.title = "Projects"
        missing.create_sheet("Tasks")
        payload = BytesIO()
        missing.save(payload)

        result = read_workbook(upload(payload.getvalue()))

        self.assertTrue(any("Instructions" in issue.message for issue in result.errors))

        result = read_workbook(upload(workbook_bytes(project_headers=("wrong",))))
        self.assertTrue(any(issue.sheet == "Projects" and issue.field == "headers" for issue in result.errors))

    def test_non_xlsx_and_oversized_uploads_are_rejected(self):
        wrong_type = read_workbook(upload(b"not excel", name="projects.xlsm"))
        oversized = read_workbook(upload(b"x" * (2 * 1024 * 1024 + 1)))

        self.assertIn(".xlsx", wrong_type.errors[0].message)
        self.assertIn("2 MiB", oversized.errors[0].message)

    def test_formula_cells_are_rejected_without_evaluation(self):
        values = [""] * len(PROJECT_COLUMNS)
        values[0] = "project-1"
        values[1] = "=1+1"

        result = read_workbook(upload(workbook_bytes(project_rows=[values])))

        self.assertTrue(any(issue.row == 2 and "formula" in issue.message.lower() for issue in result.errors))

    def test_renamed_macro_enabled_file_is_rejected(self):
        original = ZipFile(BytesIO(workbook_bytes()), "r")
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as target:
            for item in original.infolist():
                target.writestr(item, original.read(item.filename))
            target.writestr("xl/vbaProject.bin", b"macro")

        result = read_workbook(upload(output.getvalue(), name="renamed.xlsx"))

        self.assertTrue(any("macro" in issue.message.lower() for issue in result.errors))

    def test_project_and_task_row_limits_are_enforced(self):
        project_row = [""] * len(PROJECT_COLUMNS)
        task_row = [""] * len(TASK_COLUMNS)
        project_row[0] = "project-row"
        task_row[0] = "Project"

        too_many_projects = read_workbook(upload(workbook_bytes(project_rows=[project_row] * 251)))
        too_many_tasks = read_workbook(upload(workbook_bytes(task_rows=[task_row] * 2001)))

        self.assertTrue(any("250" in issue.message for issue in too_many_projects.errors))
        self.assertTrue(any("2,000" in issue.message for issue in too_many_tasks.errors))
