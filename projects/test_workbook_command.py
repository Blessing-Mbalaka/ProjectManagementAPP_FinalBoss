import tempfile
from io import BytesIO, StringIO
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import CommandError, call_command
from django.test import TestCase
from openpyxl import load_workbook

from adminpanel.models import ResearchCentre
from projects.models import Project
from projects.workbook_import.schema import PROJECT_COLUMNS
from projects.workbook_import.workbook import build_template_workbook


User = get_user_model()


class WorkbookImportCommandTests(TestCase):
    def setUp(self):
        self.centre = ResearchCentre.objects.create(name="Command Centre")
        self.manager = User.objects.create_user(username="manager1", role="manager", research_centre=self.centre)

    def write_workbook(self, **overrides):
        workbook = load_workbook(BytesIO(build_template_workbook()))
        projects = workbook["Projects"]
        projects.delete_rows(2, projects.max_row)
        row = {column: "" for column in PROJECT_COLUMNS}
        row.update({
            "import_key": "command-1",
            "project_name": "Command Project",
            "portfolio": "Command Portfolio",
            "description": "Imported by command",
            "project_type": "book",
            "project_status": "planning",
            "due_date": "2027-03-31",
            "project_manager_username": "manager1",
            "research_centre": self.centre.name,
            "publication_status": "writing",
            "lead_author_username": "missing-author",
            "external_authors": "External Writer",
            "total_chapters": "2",
            "completed_chapters": "0",
        })
        row.update(overrides)
        projects.append([row[column] for column in PROJECT_COLUMNS])
        workbook["Tasks"].delete_rows(2, workbook["Tasks"].max_row)
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        self.addCleanup(Path(handle.name).unlink, missing_ok=True)
        return handle.name

    def test_dry_run_reports_warning_and_writes_nothing(self):
        path = self.write_workbook()
        output = StringIO()

        call_command("import_project_workbook", path, dry_run=True, stdout=output)

        text = output.getvalue()
        self.assertIn("Warnings: 1", text)
        self.assertIn("missing-author", text)
        self.assertFalse(Project.objects.exists())

    def test_live_run_is_idempotent(self):
        path = self.write_workbook()

        call_command("import_project_workbook", path)
        call_command("import_project_workbook", path)

        self.assertEqual(Project.objects.count(), 1)
        self.assertEqual(Project.objects.get().portfolio, "Command Portfolio")

    def test_live_run_rejects_validation_errors(self):
        path = self.write_workbook(project_manager_username="missing-manager")

        with self.assertRaises(CommandError) as error:
            call_command("import_project_workbook", path)

        self.assertIn("Projects row 2", str(error.exception))
        self.assertFalse(Project.objects.exists())
